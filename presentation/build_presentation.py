#!/usr/bin/env python
"""
Build the GoldWatch presentation for Data Analytics and Visualization.

Run from the project root (after the data pipeline and anomaly detection):
    pip install python-pptx matplotlib
    python presentation/build_presentation.py

Outputs:
    presentation/figures/*.png            charts generated from the real data
    presentation/GoldWatch_Presentation.pptx
Dashboard screenshots (presentation/figures/dashboard_*.png) are used if present.
"""

import math
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.anomaly_detection import EVENT_DOWN, EVENT_UP, EVENT_VOLUME, detect_anomalies  # noqa: E402
from src.backtest import forward_returns, run_backtest  # noqa: E402

OUT_DIR = ROOT / "presentation"
FIG_DIR = OUT_DIR / "figures"
PPTX_PATH = OUT_DIR / "GoldWatch_Presentation.pptx"

# Colours taken from the dashboard theme (assets/style.css)
BG = "071019"
SURFACE = "0B1722"
ELEVATED = "10232D"
BORDER = "223C52"
GOLD = "F5BD4F"
TEAL = "14B8A6"
CYAN = "38BDF8"
RED = "F43F5E"
GREEN = "10B981"
TEXT = "F1F5F9"
MUTED = "94A3B8"
DIM = "64748B"

FONT = "Calibri"
MONO = "Consolas"
SUBJECT = "Data Analytics and Visualization"


# =============================================================================
# DATA & STATISTICS
# =============================================================================
def load_data():
    df = pd.read_csv(ROOT / "data" / "processed" / "gold_futures_processed.csv", parse_dates=["Date"])
    raw_rows = len(pd.read_csv(ROOT / "data" / "raw" / "gold_futures_raw.csv"))
    analysis = detect_anomalies(df)
    summary = run_backtest(analysis).set_index("Horizon_Days")
    return df, raw_rows, analysis, summary


def compute_stats(df, raw_rows, analysis, summary):
    r = analysis["Daily_Return"].dropna()
    z = analysis["Return_ZScore"].dropna()
    vol = analysis["Volume"].where(analysis["Volume"] > 0)
    years = analysis["Date"].dt.year
    counts = analysis["Event_Type"].value_counts()
    s20 = summary.loc[20]
    return {
        "raw_rows": raw_rows,
        "rows": len(df),
        "removed": raw_rows - len(df),
        "start": df["Date"].min(),
        "end": df["Date"].max(),
        "first_close": df["Close"].iloc[0],
        "last_close": df["Close"].iloc[-1],
        "kurtosis": r.kurt(),
        "share_z": (z.abs() >= 2.5).mean(),
        "normal_share": 2 * (1 - 0.5 * (1 + math.erf(2.5 / math.sqrt(2)))),
        "vol_pre": vol[years < 2020].median(),
        "vol_post": vol[years >= 2020].median(),
        "zero_vol_pre": int((analysis.loc[years < 2020, "Volume"] <= 0).sum()),
        "n_up": int(counts.get(EVENT_UP, 0)),
        "n_down": int(counts.get(EVENT_DOWN, 0)),
        "n_vol": int(counts.get(EVENT_VOLUME, 0)),
        "n_anom": int(analysis["Is_Anomaly"].sum()),
        "share_anom": analysis["Is_Anomaly"].mean(),
        "s20": s20,
        "se20": math.sqrt(s20.Baseline_Decline_Rate * (1 - s20.Baseline_Decline_Rate) / s20.Events),
    }


# =============================================================================
# CHARTS (Matplotlib, styled like the dashboard)
# =============================================================================
plt.rcParams.update({
    "figure.facecolor": "#" + SURFACE, "axes.facecolor": "#" + SURFACE, "savefig.facecolor": "#" + SURFACE,
    "axes.edgecolor": "#" + BORDER, "axes.labelcolor": "#" + MUTED, "xtick.color": "#" + MUTED,
    "ytick.color": "#" + MUTED, "text.color": "#" + TEXT, "grid.color": "#162838", "axes.grid": True,
    "font.size": 12, "axes.labelsize": 12, "legend.frameon": False, "legend.fontsize": 11,
    "axes.spines.top": False, "axes.spines.right": False, "font.family": "DejaVu Sans",
})


def save(fig, name):
    path = FIG_DIR / f"{name}.png"
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return path


def pct_axis(ax, axis="y", decimals=0):
    fmt = matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.{decimals}f}%")
    (ax.yaxis if axis == "y" else ax.xaxis).set_major_formatter(fmt)


def chart_title_spark(a):
    fig, ax = plt.subplots(figsize=(7, 4.2))
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")
    ax.plot(a["Date"], a["Close"], color="#" + GOLD, lw=2.2)
    ax.fill_between(a["Date"], a["Close"], color="#" + GOLD, alpha=0.10)
    ups = a[a["Event_Type"] == EVENT_UP]
    ax.scatter(ups["Date"], ups["Close"], s=14, color="#" + TEAL, zorder=3)
    ax.axis("off")
    path = FIG_DIR / "title_spark.png"
    fig.savefig(path, dpi=200, transparent=True, bbox_inches="tight")
    plt.close(fig)
    return path


def chart_price_history(a, st):
    fig, ax = plt.subplots(figsize=(7.2, 4.4))
    ax.plot(a["Date"], a["Close"], color="#" + GOLD, lw=1.6)
    ax.fill_between(a["Date"], a["Close"], color="#" + GOLD, alpha=0.12)
    ax.set_ylabel("Close price (USD / troy oz)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))
    ax.annotate(f"${st['first_close']:,.0f}", (a["Date"].iloc[0], st["first_close"]), xytext=(5, 18),
                textcoords="offset points", color="#" + TEXT, fontsize=11)
    ax.annotate(f"${st['last_close']:,.0f}", (a["Date"].iloc[-1], st["last_close"]), xytext=(-60, 10),
                textcoords="offset points", color="#" + TEXT, fontsize=11)
    return save(fig, "price_history")


def chart_returns_hist(a):
    r = a["Daily_Return"].dropna() * 100
    fig, ax = plt.subplots(figsize=(6.2, 4.0))
    ax.hist(r, bins=np.linspace(-7, 7, 141), density=True, color="#" + TEAL, alpha=0.85, label="Gold daily returns")
    x = np.linspace(-7, 7, 400)
    mu, sd = r.mean(), r.std()
    ax.plot(x, np.exp(-0.5 * ((x - mu) / sd) ** 2) / (sd * np.sqrt(2 * np.pi)), color="#" + GOLD, lw=2,
            label="Normal distribution")
    ax.set_yscale("log")
    ax.set_ylim(1e-4, 1.5)
    ax.set_xlabel("Daily return")
    ax.set_ylabel("Density (log scale)")
    pct_axis(ax, "x")
    ax.legend(loc="upper right")
    return save(fig, "returns_hist")


def chart_rolling_vol(a):
    vol = a["Daily_Return"].rolling(30).std() * np.sqrt(252) * 100
    fig, ax = plt.subplots(figsize=(6.2, 4.0))
    ax.plot(a["Date"], vol, color="#" + CYAN, lw=1.2)
    ax.axhline(vol.mean(), color="#" + MUTED, ls="--", lw=1)
    ax.text(a["Date"].iloc[60], vol.mean() + 1, "long-run average", color="#" + MUTED, fontsize=10)
    ax.set_ylabel("30-day volatility (annualised)")
    pct_axis(ax)
    return save(fig, "rolling_volatility")


def chart_volume_years(a):
    v = a.assign(Year=a["Date"].dt.year, V=a["Volume"].where(a["Volume"] > 0))
    med = v.groupby("Year")["V"].median()
    fig, ax = plt.subplots(figsize=(7.2, 4.4))
    colors = ["#" + TEAL if y >= 2020 else "#475569" for y in med.index]
    ax.bar(med.index, med.values, color=colors)
    ax.set_yscale("log")
    ax.axvline(2019.5, color="#" + GOLD, ls="--", lw=1.5)
    ax.text(2019.3, med.max() * 1.6, "volume used\nfrom 2020 →", color="#" + GOLD, ha="right", fontsize=11)
    ax.set_ylabel("Median daily volume (log scale)")
    ax.set_ylim(top=med.max() * 8)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    return save(fig, "volume_by_year")


def chart_zscore_window(a, date):
    i = a.index[a["Date"] == pd.Timestamp(date)][0]
    past = a["Daily_Return"].iloc[i - 30:i] * 100
    today = a["Daily_Return"].iloc[i] * 100
    mu, sd = a.loc[i, "Return_Rolling_Mean"] * 100, a.loc[i, "Return_Rolling_Std"] * 100
    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    ax.axhspan(mu - 2.5 * sd, mu + 2.5 * sd, color="#" + TEAL, alpha=0.10, label="μ ± 2.5σ (normal range)")
    ax.bar(range(-30, 0), past, color="#64748B", label="previous 30 days (window)")
    ax.bar([0], [today], color="#" + GOLD, label="today")
    ax.axhline(mu, color="#" + TEXT, ls="--", lw=1, label="mean μ")
    ax.annotate(f"Z = {a.loc[i, 'Return_ZScore']:.1f}", (0, today), xytext=(-95, -8), textcoords="offset points",
                color="#" + GOLD, fontsize=14, fontweight="bold")
    ax.set_xlabel("Trading day relative to today (t = 0)")
    ax.set_ylabel("Daily return")
    pct_axis(ax, decimals=1)
    ax.legend(loc="upper left", fontsize=10)
    ax.set_ylim(top=max(today, mu + 2.5 * sd) * 1.45)
    return save(fig, "zscore_window"), today, mu, sd, a.loc[i, "Return_ZScore"]


def chart_zscore_2013(a):
    w = a[(a["Date"] >= "2013-01-01") & (a["Date"] <= "2013-07-31")]
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(12, 4.2), sharex=True, gridspec_kw={"height_ratios": [1, 1.1]})
    ax1.plot(w["Date"], w["Close"], color="#" + GOLD, lw=1.8)
    ax1.set_ylabel("Close (USD)")
    flagged = w[w["Event_Type"] == EVENT_DOWN]
    ax1.scatter(flagged["Date"], flagged["Close"], color="#" + RED, marker="v", s=60, zorder=3)
    colors = np.where(w["Return_ZScore"] >= 0, "#" + GREEN, "#" + RED)
    ax2.bar(w["Date"], w["Return_ZScore"], color=colors, width=1.0)
    for y in (2.5, -2.5):
        ax2.axhline(y, color="#" + GOLD, ls="--", lw=1.2)
    ax2.text(w["Date"].iloc[2], 2.9, "+2.5 threshold", color="#" + GOLD, fontsize=10)
    ax2.text(w["Date"].iloc[2], -3.9, "−2.5 threshold", color="#" + GOLD, fontsize=10)
    ax2.set_ylabel("Return Z-score")
    return save(fig, "zscore_2013")


def chart_events_by_year(a):
    t = pd.crosstab(a["Date"].dt.year, a["Event_Type"]).reindex(columns=[EVENT_UP, EVENT_DOWN, EVENT_VOLUME],
                                                                fill_value=0)
    fig, ax = plt.subplots(figsize=(6.2, 3.0))
    bottom = np.zeros(len(t))
    for col, c in [(EVENT_UP, GREEN), (EVENT_DOWN, RED), (EVENT_VOLUME, CYAN)]:
        ax.bar(t.index, t[col], bottom=bottom, color="#" + c, label=col)
        bottom += t[col].values
    ax.set_ylabel("Anomalies per year")
    ax.legend(loc="upper left", ncol=3, fontsize=10)
    ax.set_ylim(top=bottom.max() * 1.35)
    return save(fig, "events_by_year")


def chart_price_anomalies(a):
    fig, ax = plt.subplots(figsize=(8.0, 4.8))
    ax.plot(a["Date"], a["Close"], color="#" + GOLD, lw=1.2, alpha=0.9, label="Gold close")
    for ev, marker, c in [(EVENT_UP, "^", GREEN), (EVENT_DOWN, "v", RED), (EVENT_VOLUME, "D", CYAN)]:
        e = a[a["Event_Type"] == ev]
        ax.scatter(e["Date"], e["Close"], marker=marker, s=26 if marker != "D" else 18, color="#" + c,
                   label=ev, zorder=3)
    ax.set_yscale("log")
    ax.set_ylabel("Close (USD, log scale)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"${v:,.0f}"))
    ax.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    ax.legend(loc="upper left")
    return save(fig, "price_anomalies")


def chart_backtest_bars(summary):
    h = summary.index.tolist()
    x = np.arange(len(h))
    fig, ax = plt.subplots(figsize=(6.6, 4.4))
    ax.bar(x - 0.2, summary["Decline_Rate"] * 100, 0.4, color="#" + GOLD, label="After an upward spike")
    ax.bar(x + 0.2, summary["Baseline_Decline_Rate"] * 100, 0.4, color="#475569", label="Any day (baseline)")
    for xi, (_, row) in zip(x, summary.iterrows()):
        label = f"p = {row.P_Value:.3f}" + (" ★" if row.Significant else "")
        ax.text(xi, max(row.Decline_Rate, row.Baseline_Decline_Rate) * 100 + 1.5, label, ha="center",
                color="#" + (GOLD if row.Significant else MUTED), fontsize=11,
                fontweight="bold" if row.Significant else "normal")
    ax.set_xticks(x, [f"{d} day" + ("s" if d > 1 else "") for d in h])
    ax.set_ylabel("Share of cases where price fell")
    ax.set_ylim(0, 72)
    pct_axis(ax)
    ax.legend(loc="upper left")
    return save(fig, "backtest_decline_rates")


def chart_event_path(a, max_h=20):
    close = a["Close"].reset_index(drop=True)
    has_z = a["Return_ZScore"].notna().values
    spikes = np.where((a["Event_Type"] == EVENT_UP).values)[0]
    spikes = spikes[spikes + max_h < len(close)]
    days = np.arange(max_h + 1)
    paths = np.array([[close[t + k] / close[t] - 1 for k in days] for t in spikes]) * 100
    mean = paths.mean(axis=0)
    se = paths.std(axis=0, ddof=1) / np.sqrt(len(spikes))
    base = np.array([forward_returns(close, k)[has_z].mean() if k else 0.0 for k in days]) * 100

    fig, ax = plt.subplots(figsize=(6.6, 4.4))
    ax.fill_between(days, mean - 1.96 * se, mean + 1.96 * se, color="#" + GOLD, alpha=0.15,
                    label="95% confidence band")
    ax.plot(days, mean, color="#" + GOLD, lw=2.4, marker="o", ms=3.5, label=f"After upward spike (n={len(spikes)})")
    ax.plot(days, base, color="#" + MUTED, lw=2, ls="--", label="Any day (baseline)")
    ax.axhline(0, color="#" + DIM, lw=1)
    ax.set_xlabel("Trading days after the spike")
    ax.set_ylabel("Average cumulative return")
    pct_axis(ax, decimals=1)
    ax.legend(loc="upper left")
    return save(fig, "event_study_path"), mean[-1], base[-1], len(spikes)


# =============================================================================
# SLIDE HELPERS
# =============================================================================
def rgb(h):
    return RGBColor.from_string(h)


def set_bullet(p, color=GOLD, indent=0.28):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def add_runs(p, content, size, color, bold, font):
    segments = content if isinstance(content, list) else [content]
    for seg in segments:
        txt, style = (seg, {}) if isinstance(seg, str) else seg
        run = p.add_run()
        run.text = txt
        f = run.font
        f.name = style.get("font", font)
        f.size = Pt(style.get("size", size))
        f.bold = style.get("bold", bold)
        f.italic = style.get("italic", False)
        f.color.rgb = rgb(style.get("color", color))
        if style.get("sub"):
            run._r.get_or_add_rPr().set("baseline", "-25000")


def text(slide, x, y, w, h, content, size=16, color=TEXT, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=6, bullets=False, bullet_color=GOLD):
    """Text box. `content` is a string or a list of paragraphs; a paragraph is a string or a list of
    segments, where a segment is a string or (string, {bold/color/size/font/italic})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        if bullets:
            set_bullet(p, bullet_color)
        add_runs(p, para, size, color, bold, font)
    return tb


def box(slide, x, y, w, h, fill=SURFACE, line=BORDER, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, line_w=1):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def card(slide, x, y, w, h, label=None, label_color=TEAL, fill=SURFACE, line=BORDER):
    box(slide, x, y, w, h, fill=fill, line=line)
    if label:
        text(slide, x + 0.22, y + 0.14, w - 0.44, 0.35, label.upper(), size=12, color=label_color, bold=True)


def chip(slide, x, y, w, h, label, fill=ELEVATED, color=TEXT, size=13, bold=True, line=BORDER):
    box(slide, x, y, w, h, fill=fill, line=line, radius=0.3)
    text(slide, x, y, w, h, label, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, after=0)


def arrow(slide, x, y, w=0.3, h=0.3, color=GOLD):
    box(slide, x, y, w, h, fill=color, line=None, kind=MSO_SHAPE.RIGHT_ARROW)


def picture(slide, path, x, y, w=None, h=None):
    """Place an image inside the (x, y, w, h) box, keeping its aspect ratio, centred."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih) if (w and h) else (w / iw if w else h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + ((w - pw) / 2 if w else 0)
    py = y + ((h - ph) / 2 if h else 0)
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def code(slide, x, y, w, h, lines, size=14):
    box(slide, x, y, w, h, fill="08131C", line=BORDER, radius=0.04)
    paras = []
    for ln in lines:
        if "#" in ln:
            body, comment = ln.split("#", 1)
            paras.append([(body, {}), ("#" + comment, {"color": DIM})])
        else:
            paras.append(ln)
    text(slide, x + 0.2, y + 0.15, w - 0.4, h - 0.3, paras, size=size, font=MONO, color=TEXT, spacing=1.05,
         after=2)


def set_cell_border(cell, color=BORDER, width=12700):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        ln = etree.Element(qn(tag), w=str(width))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
        tcPr.insert(0, ln)


def table(slide, x, y, w, rows, col_widths, row_h=0.42, size=13, highlight=None, align=None):
    """rows[0] is the header. `highlight` = set of row indexes drawn in gold."""
    highlight = highlight or set()
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                Inches(row_h * len(rows)))
    tbl = gf._element.graphic.graphicData.tbl
    tbl.tblPr.set("bandRow", "0")
    t = gf.table
    for j, cw in enumerate(col_widths):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(ELEVATED if i == 0 else SURFACE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_cell_border(cell)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (align[j] if align else PP_ALIGN.LEFT)
            color = TEAL if i == 0 else (GOLD if i in highlight else TEXT)
            add_runs(p, str(val), size if i else size - 1, color, i == 0 or i in highlight, FONT)
    return t


class Deck:
    W, H = 13.333, 7.5

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.W)
        self.prs.slide_height = Inches(self.H)
        self.count = 0

    def slide(self, title=None, tag=None, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.count += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BG)
        if title:
            box(s, 0.6, 0.55, 0.08, 0.6, fill=GOLD, line=None, kind=MSO_SHAPE.RECTANGLE)
            text(s, 0.82, 0.42, 9.6, 0.85, title, size=30, bold=True, anchor=MSO_ANCHOR.MIDDLE, after=0)
        if tag:
            text(s, 9.6, 0.58, 3.13, 0.45, tag.upper(), size=12, color=TEAL, bold=True, align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE, after=0)
        if self.count > 1:
            text(s, 0.6, 7.0, 8, 0.32, f"GoldWatch  ·  {SUBJECT}", size=10, color=DIM, after=0)
            text(s, 11.73, 7.0, 1.0, 0.32, str(self.count), size=10, color=DIM, align=PP_ALIGN.RIGHT, after=0)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


# =============================================================================
# SLIDES
# =============================================================================
def build(df, raw_rows, a, summary):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    st = compute_stats(df, raw_rows, a, summary)
    s20 = st["s20"]

    fig = {
        "spark": chart_title_spark(a),
        "price": chart_price_history(a, st),
        "hist": chart_returns_hist(a),
        "vol": chart_rolling_vol(a),
        "volume": chart_volume_years(a),
        "z2013": chart_zscore_2013(a),
        "years": chart_events_by_year(a),
        "anoms": chart_price_anomalies(a),
        "bars": chart_backtest_bars(summary),
    }
    fig["window"], brexit_ret, brexit_mu, brexit_sd, brexit_z = chart_zscore_window(a, "2016-06-24")
    fig["path"], path_spike, path_base, path_n = chart_event_path(a)

    d = Deck()

    # 1 ── Title ---------------------------------------------------------------
    s = d.slide(notes=(
        "Speaker: any member.\n"
        "Introduce the team and the project: GoldWatch detects unusual price and volume moves in gold futures "
        "using rolling statistics, and tests whether upward spikes are followed by declines."))
    picture(s, fig["spark"], 6.9, 0.9, 6.0, 4.2)
    text(s, 0.6, 1.25, 7.0, 0.45, f"{SUBJECT.upper()}  ·  COURSE PROJECT", size=14, color=TEAL, bold=True)
    text(s, 0.6, 1.75, 7.0, 1.2, "GoldWatch", size=66, color=GOLD, bold=True, after=0)
    text(s, 0.6, 3.0, 6.6, 1.2, "Detecting anomalous price spikes in gold futures with rolling statistics",
         size=24, after=0)
    text(s, 0.6, 4.25, 6.4, 0.8, "Can price and volume statistics alone flag pump-and-dump-style moves?",
         size=16, color=MUTED)
    for i, (who, role) in enumerate([("Member A", "Data Collection & Preprocessing"),
                                     ("Member B", "Anomaly Detection & Backtesting"),
                                     ("Member C", "Dashboard & Integration")]):
        x = 0.6 + i * 4.11
        card(s, x, 5.35, 3.89, 1.3, label=who)
        text(s, x + 0.22, 5.72, 3.5, 0.4, "[Name]", size=18, bold=True, after=0)
        text(s, x + 0.22, 6.12, 3.5, 0.4, role, size=13, color=MUTED, after=0)

    # 2 ── Agenda --------------------------------------------------------------
    s = d.slide("Agenda", notes="Speaker: any member. Walk through the structure of the talk in one sentence each.")
    items = [
        ("Problem & Motivation", "From crypto pump-and-dumps to gold"),
        ("Dataset & Preprocessing", "Yahoo Finance GC=F, 2000–2026"),
        ("Exploratory Data Analysis", "Fat tails, volatility clusters, volume quality"),
        ("Anomaly Detection", "Algorithm 1: rolling Z-score · Algorithm 2: scoring"),
        ("Backtesting", "Algorithm 3: event study · Algorithm 4: Z-test"),
        ("Interactive Dashboard", "Streamlit + Plotly"),
        ("Findings & Limitations", "What the data says — and what it can't"),
    ]
    for i, (t, sub) in enumerate(items):
        col, row = divmod(i, 4)
        x, y = 0.6 + col * 6.23, 1.6 + row * 1.3
        card(s, x, y, 5.9, 1.1)
        text(s, x + 0.2, y, 0.8, 1.1, f"{i + 1:02d}", size=28, color=GOLD, bold=True, anchor=MSO_ANCHOR.MIDDLE,
             after=0)
        text(s, x + 1.05, y + 0.14, 4.7, 0.45, t, size=19, bold=True, after=0)
        text(s, x + 1.05, y + 0.57, 4.7, 0.4, sub, size=13, color=MUTED, after=0)

    # 3 ── Problem -------------------------------------------------------------
    s = d.slide("Problem & Motivation", tag="Introduction", notes=(
        "Speaker: Member A or B.\n"
        "Our starting point was a crypto project idea (PumpWatch): in small coins, groups buy together, the price "
        "pumps on heavy volume, then they sell and it dumps. We asked whether the same statistical tools work on "
        "gold — a huge, liquid market — and whether upward spikes there are also followed by declines."))
    card(s, 0.6, 1.55, 5.75, 5.2, label="Inspiration: pump-and-dump")
    text(s, 0.82, 2.0, 5.3, 1.6, "In small-cap crypto, groups buy a coin together. The price “pumps” on heavy "
         "volume, then they sell and the price “dumps”.", size=16)
    for i, (lbl, c) in enumerate([("Price spike", GREEN), ("Heavy volume", CYAN), ("Crash", RED)]):
        x = 0.82 + i * 1.85
        chip(s, x, 3.55, 1.5, 0.55, lbl, color=c, size=13)
        if i < 2:
            arrow(s, x + 1.53, 3.68, 0.28, 0.28)
    text(s, 0.82, 4.4, 5.3, 2.2, [
        "PumpWatch idea: flag these spikes using only price and volume statistics.",
        [("Our twist: ", {"bold": True, "color": GOLD}), "apply the same idea to gold futures."],
    ], size=16, bullets=True)

    card(s, 6.6, 1.55, 6.13, 2.45, label="Research question", label_color=GOLD, line=GOLD)
    text(s, 6.82, 2.0, 5.7, 1.9, "Can unusual price and volume moves in gold be detected statistically — and are "
         "upward spikes followed by declines?", size=21, bold=True)
    card(s, 6.6, 4.2, 6.13, 1.65, label="Why gold?")
    text(s, 6.82, 4.58, 5.75, 1.3, [
        "Huge, liquid market: a tough test for a pattern seen in crypto",
        "26 years of free daily data, and spikes follow real news",
    ], size=15, bullets=True, after=4)
    for i, lbl in enumerate(["Detect", "Score", "Backtest", "Visualize"]):
        chip(s, 6.6 + i * 1.57, 6.1, 1.42, 0.55, lbl, color=GOLD, size=14)

    # 4 ── Dataset -------------------------------------------------------------
    s = d.slide("Dataset", tag="Data", notes=(
        "Speaker: Member A.\n"
        f"We use gold futures (ticker GC=F) from Yahoo Finance through the yfinance library: {st['raw_rows']:,} "
        f"daily rows from {st['start']:%B %Y} to {st['end']:%B %Y}. Each row has open, high, low, close and volume. "
        "Because the price rose more than tenfold, we analyse percentage returns rather than raw prices."))
    card(s, 0.6, 1.55, 4.9, 5.2, label="At a glance")
    facts = [
        ("Instrument", "COMEX gold futures (GC=F)"),
        ("Source", "Yahoo Finance via yfinance"),
        ("Period", f"{st['start']:%d %b %Y} – {st['end']:%d %b %Y}"),
        ("Frequency", "Daily trading sessions"),
        ("Rows", f"{st['raw_rows']:,} raw → {st['rows']:,} clean"),
        ("Fields", "Open, High, Low, Close, Volume"),
    ]
    for i, (k, v) in enumerate(facts):
        y = 2.05 + i * 0.75
        text(s, 0.82, y, 4.5, 0.3, k.upper(), size=11, color=DIM, bold=True, after=0)
        text(s, 0.82, y + 0.27, 4.5, 0.4, v, size=17, after=0)
    card(s, 5.75, 1.55, 6.98, 5.2, label="Gold close price, 2000–2026")
    picture(s, fig["price"], 5.9, 1.95, 6.68, 4.05)
    text(s, 5.97, 6.08, 6.55, 0.6, f"Price rose from ${st['first_close']:,.0f} to ${st['last_close']:,.0f} per "
         "ounce, so we analyse % returns, not raw prices.", size=13, color=MUTED, after=0)

    # 5 ── Workflow ------------------------------------------------------------
    s = d.slide("System Workflow", tag="Architecture", notes=(
        "Speaker: any member.\n"
        "The project is a pipeline of independent stages. Each stage writes a CSV that the next stage reads, so "
        "every step can be run and checked on its own. Member A builds the data pipeline, Member B the detection "
        "and backtest, Member C the dashboard that puts it all together."))
    stages = [
        ("SOURCE", "Yahoo Finance", "GC=F daily OHLCV via yfinance", "API download", DIM),
        ("MEMBER A", "Data Pipeline", "Clean, validate, compute returns", "gold_futures_processed.csv", TEAL),
        ("MEMBER B", "Anomaly Detection", "Rolling Z-scores, anomaly score", "gold_futures_analysis.csv", GOLD),
        ("MEMBER B", "Backtest", "Spike → forward return vs. baseline", "gold_futures_backtest.csv", GOLD),
        ("MEMBER C", "Dashboard", "Interactive charts & controls", "streamlit run app.py", CYAN),
    ]
    for i, (owner, name, desc, out, c) in enumerate(stages):
        x = 0.6 + i * 2.55
        box(s, x, 1.75, 2.15, 2.1, fill=SURFACE, line=c, line_w=1.5)
        text(s, x + 0.15, 1.9, 1.9, 0.3, owner, size=11, color=c, bold=True, after=0)
        text(s, x + 0.15, 2.22, 1.9, 0.75, name, size=18, bold=True, after=0, spacing=0.95)
        text(s, x + 0.15, 2.95, 1.9, 0.85, desc, size=13, color=MUTED, after=0)
        text(s, x, 3.95, 2.15, 0.5, out, size=11, color=c, font=MONO, align=PP_ALIGN.CENTER, after=0)
        if i < 4:
            arrow(s, x + 2.2, 2.65, 0.3, 0.3)
    text(s, 0.6, 4.6, 12.13, 0.5, "Each stage writes a CSV that the next stage reads, so every step can be run "
         "and tested on its own.", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    card(s, 0.6, 5.3, 12.13, 1.4, label="Tech stack")
    for i, tool in enumerate(["Python", "pandas", "NumPy", "yfinance", "Streamlit", "Plotly", "Matplotlib",
                              "pytest"]):
        chip(s, 0.82 + i * 1.47, 5.85, 1.32, 0.5, tool, size=13, bold=False)

    # 6 ── Preprocessing -------------------------------------------------------
    s = d.slide("Data Collection & Preprocessing", tag="Member A", notes=(
        "Speaker: Member A.\n"
        "Five steps: download, normalise, validate, handle missing values, compute features. Rows with impossible "
        f"prices (for example High below Low) are removed — {st['removed']} rows in total, almost all from early "
        "2000s settlement-price issues. We never forward-fill prices, because copying yesterday's price hides real "
        "moves and would distort the anomaly detection."))
    steps = [
        ("Download", "yfinance, ticker GC=F, full history"),
        ("Normalise", "Flatten columns, standard dates, sort, remove duplicates"),
        ("Validate", "Prices > 0, High ≥ Low, Open/Close inside High–Low, Volume ≥ 0"),
        ("Missing values", "Drop rows without a Close (no forward-fill); keep missing volume as blank"),
        ("Features", "Daily return, absolute return, intraday price range"),
    ]
    for i, (t, desc) in enumerate(steps):
        y = 1.6 + i * 1.02
        card(s, 0.6, y, 7.15, 0.85)
        chip(s, 0.78, y + 0.17, 0.52, 0.52, str(i + 1), fill=GOLD, color=BG, size=16, line=None)
        text(s, 1.5, y + 0.06, 6.1, 0.38, t, size=17, bold=True, after=0)
        text(s, 1.5, y + 0.43, 6.1, 0.38, desc, size=13, color=MUTED, after=0)
    card(s, 8.0, 1.6, 4.73, 2.35, label="Result")
    text(s, 8.22, 2.0, 4.3, 0.7, f"{st['raw_rows']:,} → {st['rows']:,}", size=34, color=GOLD, bold=True, after=0)
    text(s, 8.22, 2.75, 4.3, 1.1, [f"{st['removed']} invalid rows removed", "0 duplicate dates · 0 missing closes"],
         size=14, color=MUTED, bullets=True, after=2)
    card(s, 8.0, 4.15, 4.73, 1.15, label="Daily return")
    sub = {"sub": True}
    text(s, 8.22, 4.55, 4.3, 0.6, [["R", ("t", sub), " = Close", ("t", sub), " / Close", ("t−1", sub), " − 1"]],
         size=22, font=MONO, after=0)
    card(s, 8.0, 5.5, 4.73, 1.25, label="Why no forward-fill?", label_color=RED)
    text(s, 8.22, 5.88, 4.3, 0.85, "Copying yesterday's price creates fake 0% days and hides real moves.",
         size=13, color=MUTED, after=0)

    # 7 ── EDA: returns --------------------------------------------------------
    s = d.slide("Exploratory Analysis: Returns", tag="EDA", notes=(
        "Speaker: Member B.\n"
        "Left: the histogram of daily returns on a log scale against a normal curve. Gold has fat tails — extreme "
        f"days happen far more often than a bell curve predicts (excess kurtosis {st['kurtosis']:.1f}). Right: "
        "volatility comes in clusters (2008, 2013, 2020). So a fixed threshold over the whole history would be "
        "wrong; we need a rolling window that adapts to the current regime."))
    for i, (key, label, lines) in enumerate([
        ("hist", "Distribution of daily returns", [
            [("Fat tails: ", {"bold": True, "color": GOLD}),
             f"{st['share_z']:.1%} of days have |Z| ≥ 2.5, vs. {st['normal_share']:.1%} for a normal curve"],
        ]),
        ("vol", "30-day rolling volatility", [
            [("Volatility clusters: ", {"bold": True, "color": GOLD}),
             "calm and turbulent periods alternate → use a rolling window"],
        ]),
    ]):
        x = 0.6 + i * 6.23
        card(s, x, 1.55, 5.9, 5.2, label=label)
        picture(s, fig[key], x + 0.15, 1.95, 5.6, 3.65)
        text(s, x + 0.22, 5.72, 5.45, 0.95, lines, size=14, after=0)

    # 8 ── EDA: volume ---------------------------------------------------------
    s = d.slide("Exploratory Analysis: Volume Quality", tag="EDA", notes=(
        "Speaker: Member B.\n"
        f"Before 2020, Yahoo reports a median of only about {st['vol_pre']:,.0f} contracts a day for GC=F, with "
        f"{st['zero_vol_pre']} zero-volume days. From 2020 the median is about {st['vol_post']:,.0f} — real exchange "
        "volume. So we only compute volume Z-scores from 2020. We also found that very LOW volume days fall at "
        "month-end contract rolls, so only HIGH volume counts as an anomaly."))
    card(s, 0.6, 1.55, 7.4, 5.2, label="Median daily volume by year (Yahoo Finance GC=F)")
    picture(s, fig["volume"], 0.75, 1.95, 7.1, 4.6)
    card(s, 8.25, 1.55, 4.48, 5.2, label="What we found")
    text(s, 8.47, 2.0, 4.1, 4.6, [
        [("Before 2020: ", {"bold": True}), f"median ~{st['vol_pre']:,.0f} contracts/day and "
         f"{st['zero_vol_pre']} zero days — not real exchange volume"],
        [("From 2020: ", {"bold": True}), f"median ~{st['vol_post']:,.0f}/day — real COMEX volume"],
        [("Decision: ", {"bold": True, "color": GOLD}), "volume Z-score only from 2020"],
        [("Low-volume days ", {"bold": True}), "are contract rolls at month-end → only HIGH volume is "
         "treated as an anomaly"],
    ], size=15, bullets=True, after=10)

    # 9 ── Algorithm 1 ---------------------------------------------------------
    s = d.slide("Algorithm 1: Rolling Z-score", tag="Member B · Detection", notes=(
        "Speaker: Member B.\n"
        "The Z-score measures how unusual today is compared with the previous 30 trading days. We take the mean "
        "and standard deviation of those 30 days and see how many standard deviations today is away from the mean. "
        "Today is excluded from its own window: otherwise a spike would inflate the mean and std and hide itself, "
        "and we would be using information that wasn't available yet.\n"
        f"Example (chart): Brexit vote, 24 June 2016. Gold rose {brexit_ret:.2f}% while the previous 30 days had "
        f"mean {brexit_mu:.2f}% and std {brexit_sd:.2f}%. Z = ({brexit_ret:.2f} − {brexit_mu:.2f}) / "
        f"{brexit_sd:.2f} = {brexit_z:.1f}. Complexity is O(n) with a sliding window."))
    card(s, 0.6, 1.55, 5.95, 5.2, label="Formula")
    sub = {"sub": True}
    text(s, 0.82, 1.95, 5.5, 0.6, [["Z", ("t", sub), " = (x", ("t", sub), " − μ", ("t", sub), ") / σ",
                                    ("t", sub)]], size=28, font=MONO, color=GOLD, bold=True, after=0)
    text(s, 0.82, 2.6, 5.5, 0.5, [["μ", ("t", sub), ", σ", ("t", sub),
                                   " = mean and std of the previous 30 days"]], size=14, color=MUTED, after=0)
    code(s, 0.82, 3.1, 5.5, 1.75, [
        "for each day t:",
        "    past  = x[t-30 : t]   # today excluded",
        "    mu    = mean(past)",
        "    sigma = std(past)",
        "    Z[t]  = (x[t] - mu) / sigma",
    ], size=14)
    text(s, 0.82, 5.0, 5.5, 1.7, [
        [("No look-ahead: ", {"bold": True}), "only past data is used"],
        [("Adaptive: ", {"bold": True}), "window follows calm vs. volatile periods"],
        [("Applied to: ", {"bold": True}), "daily returns and log(volume)"],
    ], size=14, bullets=True, after=4)
    card(s, 6.8, 1.55, 5.93, 5.2, label="Worked example: Brexit vote, 24 Jun 2016")
    picture(s, fig["window"], 6.95, 1.95, 5.63, 3.7)
    text(s, 7.02, 5.75, 5.5, 0.95, f"Z = ({brexit_ret:.2f}% − {brexit_mu:.2f}%) / {brexit_sd:.2f}% = "
         f"{brexit_z:.1f}  →  far outside the normal ±2.5 range", size=14, after=0)

    # 10 ── Algorithm 1 in action ---------------------------------------------
    s = d.slide("Algorithm 1 in Action: the April 2013 Gold Crash", tag="Member B · Detection", notes=(
        "Speaker: Member B.\n"
        "Top: gold price in the first half of 2013. Bottom: the rolling Z-score of daily returns, with the ±2.5 "
        "thresholds. On 12 and 15 April 2013 gold fell sharply and the Z-score drops far below −2.5, so both days "
        "are flagged (red triangles). Note that after the crash the window contains the big moves, so σ is larger "
        "and later moves of similar size get smaller Z-scores — the window adapts."))
    card(s, 0.6, 1.55, 12.13, 4.3, label="Gold price (top) and return Z-score (bottom), Jan–Jul 2013")
    picture(s, fig["z2013"], 0.75, 1.95, 11.83, 3.8)
    for i, (head, body) in enumerate([
        ("12 & 15 Apr 2013", f"Z = {a.loc[a.Date == '2013-04-12', 'Return_ZScore'].iloc[0]:.1f} and "
                             f"{a.loc[a.Date == '2013-04-15', 'Return_ZScore'].iloc[0]:.1f}: both flagged"),
        ("Threshold ±2.5", "only a few days a year cross it"),
        ("Window adapts", "after the crash σ is larger, so later moves score lower"),
    ]):
        x = 0.6 + i * 4.11
        card(s, x, 6.0, 3.89, 0.8)
        text(s, x + 0.18, 6.05, 3.6, 0.72, [[(head + "  ", {"bold": True, "color": GOLD}), body]], size=13,
             anchor=MSO_ANCHOR.MIDDLE, after=0)

    # 11 ── Algorithm 2 --------------------------------------------------------
    s = d.slide("Algorithm 2: Anomaly Scoring & Labelling", tag="Member B · Detection", notes=(
        "Speaker: Member B.\n"
        "Step 1: cap each Z-score at 5 (twice the threshold) so one extreme day doesn't dominate, and scale to "
        "0–1. Only positive volume Z counts. Step 2: weighted average — volume counts half as much as price — "
        "times 100 gives a 0–100 score. Step 3: a day is flagged if either Z crosses 2.5. Step 4: label it.\n"
        f"Result: {st['n_anom']} anomalous days ({st['share_anom']:.1%}): {st['n_up']} upward spikes, "
        f"{st['n_down']} downward shocks, {st['n_vol']} volume spikes."))
    card(s, 0.6, 1.55, 6.55, 5.2, label="Steps")
    code(s, 0.82, 2.0, 6.11, 2.45, [
        "# 1. cap and scale to 0..1",
        "r = min(|Z_return|, 5) / 5",
        "v = min(max(Z_volume, 0), 5) / 5",
        "# 2. weighted score (w = 0.5)",
        "score = 100 * (r + 0.5*v) / 1.5",
        "# 3. flag",
        "anomaly = |Z_return| >= 2.5 or Z_volume >= 2.5",
    ], size=13)
    text(s, 0.82, 4.6, 6.11, 0.35, "4. LABEL", size=12, color=TEAL, bold=True, after=0)
    for i, (rule, lbl, c) in enumerate([("Z_return ≥ +2.5", EVENT_UP, GREEN),
                                        ("Z_return ≤ −2.5", EVENT_DOWN, RED),
                                        ("only volume ≥ 2.5", EVENT_VOLUME, CYAN)]):
        y = 4.98 + i * 0.55
        text(s, 0.82, y, 3.0, 0.45, rule, size=14, font=MONO, after=0, anchor=MSO_ANCHOR.MIDDLE)
        chip(s, 4.0, y + 0.03, 2.4, 0.4, lbl, color=c, size=13)
    for i, (n, lbl, c) in enumerate([(st["n_anom"], f"anomalous days ({st['share_anom']:.1%})", GOLD),
                                     (st["n_up"], "upward spikes", GREEN),
                                     (st["n_down"], "downward shocks", RED),
                                     (st["n_vol"], "volume spikes", CYAN)]):
        x = 7.4 + i * 1.355
        card(s, x, 1.55, 1.24, 1.5)
        text(s, x, 1.68, 1.24, 0.6, str(n), size=26, color=c, bold=True, align=PP_ALIGN.CENTER, after=0)
        text(s, x + 0.05, 2.3, 1.14, 0.7, lbl, size=11, color=MUTED, align=PP_ALIGN.CENTER, after=0, spacing=0.95)
    card(s, 7.4, 3.25, 5.33, 3.5, label="Anomalies per year")
    picture(s, fig["years"], 7.5, 3.65, 5.13, 3.0)

    # 12 ── Detected anomalies ------------------------------------------------
    s = d.slide("Detected Anomalies Match Real Events", tag="Results", notes=(
        "Speaker: Member B.\n"
        "Left: the whole price history on a log scale with every flagged day. Right: some of the strongest "
        "anomalies and the news behind them. The algorithm knows nothing about the news — it only sees numbers — "
        "yet the biggest flags line up with 9/11, the 2013 crash, the Brexit vote, the Fed's 2019 turn and the "
        "Covid vaccine announcement. This is a good sanity check that the detector finds real events."))
    card(s, 0.6, 1.55, 7.35, 5.2, label="Gold price with flagged days (log scale)")
    picture(s, fig["anoms"], 0.75, 1.95, 7.05, 4.65)
    events = [("2001-09-14", "Markets reopen after 9/11"), ("2013-04-15", "April 2013 gold crash"),
              ("2016-06-24", "Brexit referendum result"), ("2019-06-20", "Fed signals rate cuts"),
              ("2020-11-09", "Covid vaccine announced")]
    rows = [["Date", "Move", "Z", "What happened"]]
    for date, what in events:
        row = a[a["Date"] == date].iloc[0]
        rows.append([pd.Timestamp(date).strftime("%d %b %Y"), f"{row.Daily_Return:+.1%}",
                     f"{row.Return_ZScore:+.1f}", what])
    table(s, 8.2, 1.55, 4.53, rows, [1.2, 0.72, 0.6, 2.01], row_h=0.62, size=12,
          align=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.LEFT])
    text(s, 8.2, 5.45, 4.53, 1.3, "The detector only sees numbers, yet its strongest flags line up with major "
         "news events.", size=14, color=MUTED, after=0)

    # 13 ── Algorithm 3 --------------------------------------------------------
    s = d.slide("Algorithm 3: Event-Study Backtest", tag="Member B · Backtest", notes=(
        "Speaker: Member B.\n"
        "For each upward spike on day t we 'buy' at the close of day t and look at the close h trading days later. "
        "The forward return is exit divided by entry minus one; if it's negative the spike was followed by a "
        "decline. We do this for h = 1, 5, 10 and 20 days.\n"
        "The key idea is the baseline: gold falls over a 5-day period about 44% of the time anyway. So we compare "
        "the decline rate after spikes with the decline rate over all days. The difference is the lift."))
    days = [("t − 1", SURFACE, TEXT), ("t", GOLD, BG), ("t + 1", SURFACE, TEXT), ("…", SURFACE, MUTED),
            ("t + h", CYAN, BG)]
    for i, (lbl, fill, color) in enumerate(days):
        x = 1.1 + i * 2.35
        box(s, x, 1.65, 1.6, 0.75, fill=fill, line=BORDER, radius=0.15)
        text(s, x, 1.65, 1.6, 0.75, lbl, size=20, bold=True, color=color, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, after=0)
        if i < 4:
            arrow(s, x + 1.75, 1.88, 0.45, 0.3, color=BORDER)
    text(s, 2.6, 2.48, 3.2, 0.7, [[("Spike day", {"bold": True, "color": GOLD})], "Entry = Close[t]"], size=14,
         align=PP_ALIGN.CENTER, after=0)
    text(s, 9.65, 2.48, 3.2, 0.7, [[("h days later", {"bold": True, "color": CYAN})], "Exit = Close[t+h]"],
         size=14, align=PP_ALIGN.CENTER, after=0)
    card(s, 0.6, 3.45, 6.55, 3.3, label="Algorithm")
    code(s, 0.82, 3.9, 6.11, 2.65, [
        "for each upward spike day t:",
        "    fwd      = Close[t+h] / Close[t] - 1",
        "    declined = fwd < 0",
        "",
        "decline_rate  = mean(declined)   # spikes",
        "baseline_rate = mean(fwd < 0)    # all days",
        "lift          = decline_rate - baseline_rate",
    ], size=13)
    card(s, 7.4, 3.45, 5.33, 3.3, label="Why compare with a baseline?", label_color=GOLD)
    text(s, 7.62, 3.9, 4.9, 2.8, [
        "Gold falls over 5 days about 44% of the time anyway.",
        "So “45% of spikes were followed by a drop” means nothing on its own.",
        [("Only the lift ", {"bold": True, "color": GOLD}), "over the baseline shows a real effect."],
        "Horizons tested: h = 1, 5, 10, 20 days.",
    ], size=14, bullets=True, after=6)

    # 14 ── Algorithm 4 --------------------------------------------------------
    s = d.slide("Algorithm 4: One-Proportion Z-test", tag="Member B · Backtest", notes=(
        "Speaker: Member B.\n"
        "The Z-test tells us whether the difference between the decline rate after spikes and the baseline is "
        "real or could be luck. H0 says spikes are no different from a normal day; H1 says they are followed by "
        "drops more often. We compute how many standard errors the observed rate is above the baseline and convert "
        "that to a p-value with the normal distribution. If p < 0.05 we reject H0.\n"
        f"Worked example with our 20-day result: {int(s20.Declines)} of {int(s20.Events)} spikes declined "
        f"({s20.Decline_Rate:.3f}) vs a baseline of {s20.Baseline_Decline_Rate:.3f}: z = {s20.Z_Stat:.2f}, "
        f"p = {s20.P_Value:.3f}, so it is significant."))
    card(s, 0.6, 1.55, 6.0, 1.15, label="H0  (null hypothesis)", label_color=MUTED)
    text(s, 0.82, 1.95, 5.6, 0.7, "Spikes are followed by drops no more often than a normal day", size=15,
         after=0)
    card(s, 0.6, 2.85, 6.0, 1.15, label="H1  (alternative)", label_color=GOLD)
    text(s, 0.82, 3.25, 5.6, 0.7, "Spikes are followed by drops MORE often", size=15, after=0)
    card(s, 0.6, 4.15, 6.0, 2.6, label="Test statistic")
    text(s, 0.82, 4.6, 5.6, 1.4, ["z = (p_spike − p_base) / SE", "SE = √( p_base(1 − p_base) / n )",
                                  "p-value = 1 − Φ(z)"], size=17, font=MONO, color=GOLD, after=4)
    text(s, 0.82, 5.85, 5.6, 0.8, "p_spike = decline rate after spikes, p_base = baseline, n = number of "
         "spikes, Φ = normal CDF. Reject H0 if p < 0.05.", size=13, color=MUTED, after=0)
    card(s, 6.85, 1.55, 5.88, 5.2, label="Worked example: 20-day horizon", label_color=GOLD, line=GOLD)
    steps = [
        ("n", f"{int(s20.Events)} spikes, {int(s20.Declines)} declined"),
        ("p_spike", f"{int(s20.Declines)} / {int(s20.Events)} = {s20.Decline_Rate:.3f}"),
        ("p_base", f"baseline = {s20.Baseline_Decline_Rate:.3f}"),
        ("SE", f"√({s20.Baseline_Decline_Rate:.3f} × {1 - s20.Baseline_Decline_Rate:.3f} / {int(s20.Events)}) "
               f"= {st['se20']:.4f}"),
        ("z", f"({s20.Decline_Rate:.3f} − {s20.Baseline_Decline_Rate:.3f}) / {st['se20']:.4f} = {s20.Z_Stat:.2f}"),
        ("p", f"1 − Φ({s20.Z_Stat:.2f}) = {s20.P_Value:.3f}"),
    ]
    for i, (k, v) in enumerate(steps):
        y = 2.05 + i * 0.6
        text(s, 7.07, y, 1.15, 0.5, k, size=15, font=MONO, color=TEAL, bold=True, after=0)
        text(s, 8.22, y, 4.4, 0.5, v, size=15, font=MONO, after=0)
    box(s, 7.07, 5.75, 5.44, 0.75, fill=ELEVATED, line=GOLD, radius=0.12)
    text(s, 7.07, 5.75, 5.44, 0.75, f"p = {s20.P_Value:.3f} < 0.05  →  reject H0", size=18, color=GOLD,
         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)

    # 15 ── Backtest results ---------------------------------------------------
    s = d.slide("Backtest Results", tag="Results", notes=(
        "Speaker: Member B.\n"
        "Gold bars: share of upward spikes followed by a decline. Grey bars: the same for any day. After 1, 5 and "
        "10 days the two are almost equal and the p-values are large — no short-term dump. Only at 20 days do "
        f"spikes decline clearly more often ({s20.Decline_Rate:.1%} vs {s20.Baseline_Decline_Rate:.1%}, "
        f"p = {s20.P_Value:.3f})."))
    card(s, 0.6, 1.55, 6.6, 5.2, label="How often did the price fall after a spike?")
    picture(s, fig["bars"], 0.75, 1.95, 6.3, 4.7)
    rows = [["Horizon", "After spike", "Baseline", "Lift", "p-value"]]
    for h, row in summary.iterrows():
        rows.append([f"{h} day" + ("s" if h > 1 else ""), f"{row.Decline_Rate:.1%}",
                     f"{row.Baseline_Decline_Rate:.1%}", f"{row.Lift:+.1%}", f"{row.P_Value:.3f}"])
    sig_rows = {i + 1 for i, v in enumerate(summary["Significant"]) if v}
    table(s, 7.45, 1.55, 5.28, rows, [1.05, 1.13, 1.05, 0.95, 1.1], row_h=0.5, size=14, highlight=sig_rows,
          align=[PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 4)
    card(s, 7.45, 4.3, 5.28, 2.45, label="Takeaway", label_color=GOLD)
    text(s, 7.67, 4.7, 4.85, 2.0, [
        "1–10 days: decline rate ≈ baseline → no short-term dump",
        f"20 days: {s20.Decline_Rate:.1%} vs. {s20.Baseline_Decline_Rate:.1%}, significant (p = {s20.P_Value:.3f})",
    ], size=15, bullets=True, after=8)

    # 16 ── Event-study path ---------------------------------------------------
    s = d.slide("What Happens After a Spike?", tag="Results", notes=(
        "Speaker: Member B.\n"
        f"This chart averages the price path over the 20 trading days after each of the {path_n} upward spikes, "
        "compared with the average path after any day. Spikes do not crash afterwards: the average stays slightly "
        f"positive. But it grows much more slowly than normal: +{path_spike:.2f}% after 20 days vs "
        f"+{path_base:.2f}% for a normal day. The band shows the 95% confidence interval. So gold spikes look like "
        "news-driven repricing that is partly given back over a month, not a pump-and-dump."))
    card(s, 0.6, 1.55, 7.2, 5.2, label="Average cumulative return after an upward spike")
    picture(s, fig["path"], 0.75, 1.95, 6.9, 4.7)
    card(s, 8.05, 1.55, 4.68, 5.2, label="Interpretation", label_color=GOLD)
    text(s, 8.27, 2.0, 4.25, 4.6, [
        [("No crash: ", {"bold": True}), "the average path stays positive after a spike"],
        [("Slower growth: ", {"bold": True}), f"+{path_spike:.2f}% after 20 days vs. +{path_base:.2f}% "
         "for a normal day"],
        [("Meaning: ", {"bold": True, "color": GOLD}), "spikes are news-driven moves that are partly "
         "given back over about a month"],
    ], size=15, bullets=True, after=12)

    # 17 ── Dashboard ----------------------------------------------------------
    s = d.slide("Interactive Dashboard", tag="Member C", notes=(
        "Speaker: Member C.\n"
        "The Streamlit dashboard loads the processed data and Member B's analysis and backtest files. The sidebar "
        "lets the user pick the date range, chart type and the model parameters — window, threshold and volume "
        "weight — and every chart updates immediately. There are four tabs: overview, anomaly signals, the "
        "event-study backtest and a data explorer with CSV export. Run it with: streamlit run app.py"))
    shot = FIG_DIR / "dashboard_overview.png"
    card(s, 0.6, 1.55, 8.3, 5.2, label="Overview tab: price, anomaly markers and volume")
    if shot.exists():
        picture(s, shot, 0.75, 1.95, 8.0, 4.65)
    else:
        text(s, 0.75, 3.8, 8.0, 0.6, "[Insert dashboard screenshot]", size=18, color=DIM, align=PP_ALIGN.CENTER)
    card(s, 9.15, 1.55, 3.58, 5.2, label="Features")
    text(s, 9.37, 2.0, 3.2, 4.65, [
        "Date range filter",
        "Area, candlestick or line price chart",
        "Sliders: window, threshold, volume weight",
        "4 tabs: Overview, Signals, Backtest, Data",
        "Ranked anomaly tables",
        "CSV export",
    ], size=14, bullets=True, after=8)

    # 18 ── Visualisation choices ---------------------------------------------
    s = d.slide("Visualisation Design Choices", tag="Visualization", notes=(
        "Speaker: Member C.\n"
        "Every chart was chosen to answer a specific question. Price with markers shows when anomalies happen; "
        "Z-score lines with threshold bands show the detection rule itself; the forward-return histogram shows "
        "the full distribution of outcomes rather than only an average; KPI cards give the headline numbers. The "
        "dashboard uses Plotly for interactivity (hover, zoom); the static charts in this deck use Matplotlib."))
    shot = FIG_DIR / "dashboard_backtest.png"
    card(s, 0.6, 1.55, 6.15, 5.2, label="Backtest tab")
    if shot.exists():
        picture(s, shot, 0.75, 1.95, 5.85, 3.2)
    text(s, 0.82, 5.3, 5.7, 1.35, "KPI cards compare the decline rate after spikes with the market baseline; "
         "the histogram shows every outcome, and a slider changes what counts as a decline.", size=13,
         color=MUTED, after=0)
    rows = [["Chart", "Question it answers"],
            ["Price + anomaly markers", "When do anomalies happen?"],
            ["Z-score lines + ±2.5σ band", "Why was this day flagged?"],
            ["Volume bars", "Was the move on heavy trading?"],
            ["Histogram of forward returns", "What usually happens next?"],
            ["Bar chart vs. baseline", "Is it better than chance?"],
            ["KPI cards & ranked tables", "Headline numbers, drill-down"]]
    table(s, 7.0, 1.55, 5.73, rows, [2.85, 2.88], row_h=0.6, size=13)
    text(s, 7.0, 5.95, 5.73, 0.8, "Interactive: Plotly (hover, zoom) · Static charts: Matplotlib", size=13,
         color=MUTED, after=0)

    # 19 ── Conclusions --------------------------------------------------------
    s = d.slide("Conclusions", tag="Summary", notes=(
        "Speaker: any member.\n"
        "Four takeaways: the detector works and finds real events; there is no short-term dump in gold; there is "
        "a slow partial reversal over about a month; and so the answer to our research question is that gold does "
        "not show crypto-style pump-and-dump behaviour."))
    findings = [
        ("Detection works", f"{st['n_anom']} anomalous days ({st['share_anom']:.1%}); the strongest match real news "
                            "(9/11, 2013 crash, Brexit, Covid vaccine).", TEAL),
        ("No short-term dump", "1–10 days after a spike, gold falls about as often as on any day (p > 0.2).", CYAN),
        ("Slow partial reversal", f"After 20 days: {s20.Decline_Rate:.1%} vs. {s20.Baseline_Decline_Rate:.1%} "
                                  f"declines (p = {s20.P_Value:.3f}).", GREEN),
        ("Answer to our question", "Gold shows no crypto-style pump-and-dump: spikes are news-driven and partly "
                                   "fade over a month.", GOLD),
    ]
    for i, (head, body, c) in enumerate(findings):
        col, row = i % 2, i // 2
        x, y = 0.6 + col * 6.23, 1.6 + row * 2.6
        card(s, x, y, 5.9, 2.35, line=c)
        text(s, x + 0.25, y + 0.2, 5.4, 0.5, head, size=21, color=c, bold=True, after=0)
        text(s, x + 0.25, y + 0.8, 5.4, 1.45, body, size=16, after=0)

    # 20 ── Limitations & future work -----------------------------------------
    s = d.slide("Limitations & Future Work", tag="Summary", notes=(
        "Speaker: any member.\n"
        "Limitations: spikes cluster in time, so their outcome windows overlap and the events are not fully "
        "independent — the p-value is somewhat optimistic. We tested four horizons; with a Bonferroni correction "
        "(0.05 / 4 = 0.0125) the 20-day result still passes. Volume is only usable from 2020. Cleaning removed a "
        "few rows, so some returns span several days. And this is not a trading strategy — no costs are modelled."))
    card(s, 0.6, 1.55, 5.9, 5.2, label="Limitations", label_color=RED)
    text(s, 0.82, 2.0, 5.45, 4.6, [
        "Spikes cluster in time → overlapping windows, p-value somewhat optimistic",
        "4 horizons tested → Bonferroni level 0.0125 (20-day result still passes)",
        "Volume only usable from 2020",
        "Removed rows make a few returns span several days",
        "No trading costs: not a trading strategy",
    ], size=15, bullets=True, bullet_color=RED, after=10)
    card(s, 6.83, 1.55, 5.9, 5.2, label="Future work", label_color=GREEN)
    text(s, 7.05, 2.0, 5.45, 4.6, [
        "Robust Z-score (median / MAD) and comparison with the standard Z-score",
        "ARIMA forecast errors and seasonal decomposition residuals as extra signals",
        "Extend to silver and other metals",
        "Intraday data for faster detection",
        "Regression of forward return on spike size",
    ], size=15, bullets=True, bullet_color=GREEN, after=10)

    # 21 ── Thank you ----------------------------------------------------------
    s = d.slide(notes="Speaker: any member. Thank the audience and invite questions.")
    picture(s, fig["spark"], 6.9, 0.9, 6.0, 4.2)
    text(s, 0.6, 2.0, 7.0, 1.2, "Thank you", size=60, color=GOLD, bold=True, after=0)
    text(s, 0.6, 3.25, 6.5, 0.7, "Questions?", size=28, after=0)
    card(s, 0.6, 5.0, 12.13, 1.7, label="References & tools")
    text(s, 0.82, 5.42, 11.7, 1.2, [
        "Data: Yahoo Finance, COMEX gold futures (GC=F), via the yfinance Python library",
        "Analysis: pandas, NumPy · Dashboard: Streamlit, Plotly · Charts: Matplotlib · Tests: pytest",
    ], size=14, color=MUTED, bullets=True, after=4)

    d.save(PPTX_PATH)
    return d.count


def main():
    df, raw_rows, analysis, summary = load_data()
    n = build(df, raw_rows, analysis, summary)
    print(f"Saved {PPTX_PATH} ({n} slides) and charts in {FIG_DIR}")


if __name__ == "__main__":
    main()
